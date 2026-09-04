#!/usr/bin/env python3
"""ppt.py — PowerPoint 技能统一入口 CLI（可移植）。

依赖：python-pptx（pip install python-pptx）。
默认输出 JSON；read 支持 --md。

命令：
  new    file.pptx [--title 标题]            新建演示文稿
  info   file.pptx                           演示文稿元信息
  read   file.pptx [--md]                    读取每页标题与文本
  search file.pptx 关键词                    搜索内容
  build  file.pptx [--data JSON]             按 JSON 结构生成 PPT
"""
import argparse
import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
except ImportError:
    print("缺少依赖 python-pptx，请先执行：pip3 install python-pptx", file=sys.stderr)
    sys.exit(2)


def _title_of(slide):
    if slide.shapes.title is not None:
        return slide.shapes.title.text
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.TITLE:
            return sh.text_frame.text
    return None


def _texts_of(slide):
    """返回每页非标题的文本框内容。"""
    out = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.is_placeholder:
            ph = sh.placeholder_format.type
            if ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                continue
        t = sh.text_frame.text.strip()
        if t:
            out.append(t)
    return out


# ---------------- 子命令 ----------------
def cmd_new(args):
    prs = Presentation()
    if args.title:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        if slide.shapes.title is not None:
            slide.shapes.title.text = args.title
    prs.save(args.file)
    print(json.dumps({"file": args.file, "title": args.title}, ensure_ascii=False))


def cmd_info(args):
    prs = Presentation(args.file)
    detail = []
    for i, slide in enumerate(prs.slides, 1):
        detail.append({"slide": i, "title": _title_of(slide), "shapes": len(slide.shapes)})
    print(json.dumps({
        "file": args.file,
        "slide_count": len(prs.slides),
        "width_emu": prs.slide_width,
        "height_emu": prs.slide_height,
        "slides": detail,
    }, ensure_ascii=False, indent=2))


def cmd_read(args):
    prs = Presentation(args.file)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        slides.append({"slide": i, "title": _title_of(slide), "texts": _texts_of(slide)})
    if args.md:
        for s in slides:
            print("## " + (s["title"] or f"第{s['slide']}页"))
            for t in s["texts"]:
                for line in t.split("\n"):
                    print("- " + line)
    else:
        print(json.dumps({"file": args.file, "slides": slides}, ensure_ascii=False, indent=2))


def cmd_search(args):
    prs = Presentation(args.file)
    needle = args.value.lower()
    hits = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text
                if needle in t.lower():
                    hits.append({"slide": i, "text": t})
    print(json.dumps({"count": len(hits), "hits": hits}, ensure_ascii=False, indent=2))


def _read_structure(args):
    if args.data:
        data = json.loads(args.data) if isinstance(args.data, str) else args.data
    else:
        data = json.load(sys.stdin)
    if not isinstance(data, list):
        raise SystemExit("PPT 结构必须是 JSON 列表（每项 {'title','bullets'[, 'layout']}）")
    return data


def cmd_build(args):
    data = _read_structure(args)
    prs = Presentation()
    for idx, item in enumerate(data):
        title = item.get("title", "")
        bullets = item.get("bullets", [])
        layout_idx = item.get("layout")
        if layout_idx is not None:
            layout = prs.slide_layouts[int(layout_idx)]
        else:
            layout = prs.slide_layouts[0] if idx == 0 else prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title is not None and title:
            slide.shapes.title.text = title
        if bullets and len(slide.placeholders) > 1:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.clear()
            for bi, b in enumerate(bullets):
                p = tf.paragraphs[0] if bi == 0 else tf.add_paragraph()
                p.text = str(b)
    prs.save(args.file)
    print(json.dumps({"file": args.file, "slides": len(data)}, ensure_ascii=False))


def main():
    p = argparse.ArgumentParser(prog="ppt.py", description="PowerPoint(.pptx) 技能统一入口 CLI")
    sub = p.add_subparsers(dest="cmd", required=True)

    sp = sub.add_parser("new", help="新建演示文稿")
    sp.add_argument("file")
    sp.add_argument("--title")
    sp.set_defaults(func=cmd_new)

    sp = sub.add_parser("info", help="演示文稿元信息")
    sp.add_argument("file")
    sp.set_defaults(func=cmd_info)

    sp = sub.add_parser("read", help="读取每页标题与文本")
    sp.add_argument("file")
    sp.add_argument("--md", action="store_true", help="输出 Markdown")
    sp.set_defaults(func=cmd_read)

    sp = sub.add_parser("search", help="搜索内容")
    sp.add_argument("file")
    sp.add_argument("value")
    sp.set_defaults(func=cmd_search)

    sp = sub.add_parser("build", help="按 JSON 结构生成 PPT")
    sp.add_argument("file")
    sp.add_argument("--data", help="JSON 结构字符串；缺省从 stdin 读取")
    sp.set_defaults(func=cmd_build)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()

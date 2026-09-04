#!/usr/bin/env python3
"""xlsx.py — Excel 技能统一入口 CLI（可移植）。

用法见 SKILL.md，或运行 `python3 xlsx.py --help`。
依赖：openpyxl（pip install openpyxl）。
默认输出 JSON，便于程序化解析。
"""
import argparse
import csv
import json
import re
import sys
from datetime import date, datetime, time

try:
    import openpyxl
    from openpyxl.utils import get_column_letter, column_index_from_string
except ImportError:
    print("缺少依赖 openpyxl，请先执行：pip3 install openpyxl", file=sys.stderr)
    sys.exit(2)


# --------------------------------------------------------------------------
# 工具函数
# --------------------------------------------------------------------------
def _jsonable(v):
    """把单元格值转为可 JSON 序列化的对象。"""
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if v is None:
        return None
    try:
        json.dumps(v)
        return v
    except (TypeError, ValueError):
        return str(v)


def load_wb(path):
    return openpyxl.load_workbook(path, data_only=False)


def _parse_range(rng):
    """把 'A1:C10' 解析为 (r1, c1, r2, c2)。"""
    rng = rng.replace("$", "")
    if ":" not in rng:
        rng = f"{rng}:{rng}"
    a, b = rng.split(":")
    ma = re.match(r"([A-Za-z]+)(\d*)", a.strip())
    mb = re.match(r"([A-Za-z]+)(\d*)", b.strip())
    if not ma or not mb:
        raise ValueError(f"非法区域: {rng}")
    c1 = column_index_from_string(ma.group(1))
    r1 = int(ma.group(2) or 1)
    c2 = column_index_from_string(mb.group(1))
    r2 = int(mb.group(2) or 1048576)
    return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


def _read_rows(ws, r1, c1, r2, c2):
    rows = []
    for r in range(r1, r2 + 1):
        rows.append([_jsonable(ws.cell(row=r, column=c).value) for c in range(c1, c2 + 1)])
    return rows


def _print_md(rows):
    if not rows:
        print("(空)")
        return
    width = max(len(r) for r in rows)
    rows = [list(r) + [""] * (width - len(r)) for r in rows]

    def esc(v):
        if v is None:
            return ""
        return str(v).replace("|", "\\|").replace("\n", " ")

    print("| " + " | ".join(esc(v) for v in rows[0]) + " |")
    print("| " + " | ".join(["---"] * width) + " |")
    for row in rows[1:]:
        print("| " + " | ".join(esc(v) for v in row) + " |")


# --------------------------------------------------------------------------
# 子命令处理
# --------------------------------------------------------------------------
def cmd_info(args):
    wb = load_wb(args.file)
    sheets = [{
        "name": ws.title,
        "rows": ws.max_row,
        "cols": ws.max_column,
        "dimensions": ws.dimensions,
    } for ws in wb.worksheets]
    print(json.dumps({"file": args.file, "sheet_count": len(sheets), "sheets": sheets},
                     ensure_ascii=False, indent=2))


def cmd_sheets(args):
    wb = load_wb(args.file)
    print(json.dumps(wb.sheetnames, ensure_ascii=False))


def cmd_read(args):
    wb = load_wb(args.file)
    ws = wb[args.sheet] if args.sheet else wb.active
    if args.range:
        r1, c1, r2, c2 = _parse_range(args.range)
        r2 = min(max(r1, r2), max(ws.max_row, 1))
        c2 = min(max(c1, c2), max(ws.max_column, 1))
    else:
        r1, c1, r2, c2 = 1, 1, ws.max_row, ws.max_column
    rows = _read_rows(ws, r1, c1, r2, c2)
    if args.md:
        _print_md(rows)
    elif args.csv:
        w = csv.writer(sys.stdout)
        for row in rows:
            w.writerow(["" if v is None else v for v in row])
    else:
        rng = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
        print(json.dumps({"sheet": ws.title, "range": rng, "rows": rows},
                         ensure_ascii=False, indent=2))


def cmd_search(args):
    wb = load_wb(args.file)
    needle = args.value.lower()
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if needle in str(cell.value).lower():
                    hits.append({"sheet": ws.title, "cell": cell.coordinate,
                                 "value": _jsonable(cell.value)})
    print(json.dumps({"count": len(hits), "hits": hits}, ensure_ascii=False, indent=2))


def _read_data(args):
    if args.data:
        data = json.loads(args.data) if isinstance(args.data, str) else args.data
    else:
        try:
            data = json.load(sys.stdin)
        except json.JSONDecodeError as e:
            raise SystemExit(f"stdin 不是合法 JSON 二维数组: {e}")
    if not isinstance(data, list):
        raise SystemExit("数据必须是二维数组（list of list）")
    return data


def cmd_write(args):
    data = _read_data(args)
    try:
        wb = openpyxl.load_workbook(args.file)
    except FileNotFoundError:
        wb = openpyxl.Workbook()
    if args.sheet in wb.sheetnames:
        ws = wb[args.sheet]
        if args.overwrite:
            wb.remove(ws)
            ws = wb.create_sheet(args.sheet, 0)
        else:
            start_row = ws.max_row + 1
            for i, row in enumerate(data):
                for j, val in enumerate(row):
                    ws.cell(row=start_row + i, column=j + 1, value=val)
            wb.save(args.file)
            print(json.dumps({"file": args.file, "sheet": args.sheet, "appended": True,
                              "start_row": start_row, "rows": len(data)},
                             ensure_ascii=False))
            return
    else:
        ws = wb.create_sheet(args.sheet)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            ws.cell(row=i + 1, column=j + 1, value=val)
    wb.save(args.file)
    cols = max((len(r) for r in data), default=0)
    print(json.dumps({"file": args.file, "sheet": args.sheet, "rows": len(data),
                      "cols": cols}, ensure_ascii=False))


def cmd_new(args):
    wb = openpyxl.Workbook()
    wb.active.title = args.sheet
    wb.save(args.file)
    print(json.dumps({"file": args.file, "sheet": args.sheet}, ensure_ascii=False))


# --------------------------------------------------------------------------
# 参数解析
# --------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(prog="xlsx.py", description="Excel 技能统一入口 CLI")
    sub = p.add_subparsers(dest="cmd", required=True)

    sp = sub.add_parser("info", help="查看工作簿元信息（sheet 列表及行列数）")
    sp.add_argument("file")
    sp.set_defaults(func=cmd_info)

    sp = sub.add_parser("sheets", help="仅列出 sheet 名称")
    sp.add_argument("file")
    sp.set_defaults(func=cmd_sheets)

    sp = sub.add_parser("read", help="读取单元格")
    sp.add_argument("file")
    sp.add_argument("sheet", nargs="?")
    sp.add_argument("--range", help="如 A1:C10")
    sp.add_argument("--md", action="store_true", help="输出 Markdown 表格")
    sp.add_argument("--csv", action="store_true", help="输出 CSV")
    sp.set_defaults(func=cmd_read)

    sp = sub.add_parser("search", help="搜索单元格内容（不区分大小写）")
    sp.add_argument("file")
    sp.add_argument("value")
    sp.set_defaults(func=cmd_search)

    sp = sub.add_parser("write", help="写入 JSON 二维数组（stdin 或 --data）")
    sp.add_argument("file")
    sp.add_argument("sheet", nargs="?", default="Sheet1")
    sp.add_argument("--data", help="JSON 二维数组字符串")
    sp.add_argument("--overwrite", action="store_true", help="覆盖同名 sheet（默认追加）")
    sp.set_defaults(func=cmd_write)

    sp = sub.add_parser("new", help="新建空白工作簿")
    sp.add_argument("file")
    sp.add_argument("sheet", nargs="?", default="Sheet1")
    sp.set_defaults(func=cmd_new)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()

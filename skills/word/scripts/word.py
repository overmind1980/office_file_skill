#!/usr/bin/env python3
"""word.py — Word 技能统一入口 CLI（可移植）。

依赖：python-docx（pip install python-docx）。
默认输出 JSON，便于程序化解析；read 支持 --md 输出 Markdown。

命令：
  new    file.docx [--title 标题]            新建文档
  info   file.docx                           文档元信息
  read   file.docx [--md] [--max N]          读取段落/表格
  search file.docx 关键词                     搜索内容
  write  file.docx [--data JSON]             按 JSON 结构写入(新建或追加)
"""
import argparse
import json
import os
import sys

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    print("缺少依赖 python-docx，请先执行：pip3 install python-docx", file=sys.stderr)
    sys.exit(2)


def _jsonable(v):
    if v is None:
        return None
    try:
        json.dumps(v)
        return v
    except (TypeError, ValueError):
        return str(v)


def iter_block_items(doc):
    """按文档真实顺序产出段落与表格。"""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _heading_level(p):
    name = (p.style.name or "").lower() if p.style else ""
    if name == "title":
        return 1
    if name.startswith("heading"):
        rest = name[len("heading"):].strip()
        if rest.isdigit():
            return max(1, min(6, int(rest)))
        return 1
    return None


def parse_doc(doc):
    blocks = []
    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            text = item.text
            lvl = _heading_level(item)
            if text.strip() == "" and lvl is None:
                continue
            if lvl is not None:
                blocks.append({"type": "heading", "level": lvl, "text": text})
            else:
                blocks.append({"type": "paragraph", "text": text})
        else:
            rows = [[_jsonable(cell.text) for cell in row.cells] for row in item.rows]
            blocks.append({"type": "table", "rows": rows})
    return blocks


def _blocks_to_md(blocks):
    out = []
    for b in blocks:
        if b["type"] == "heading":
            out.append("#" * max(1, min(b["level"], 6)) + " " + b["text"])
        elif b["type"] == "paragraph":
            out.append(b["text"])
        else:  # table
            rows = b["rows"]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            norm = [list(r) + [""] * (width - len(r)) for r in rows]

            def esc(x):
                return ("" if x is None else str(x)).replace("|", "\\|").replace("\n", " ")

            out.append("| " + " | ".join(esc(x) for x in norm[0]) + " |")
            out.append("| " + " | ".join(["---"] * width) + " |")
            for row in norm[1:]:
                out.append("| " + " | ".join(esc(x) for x in row) + " |")
    return "\n".join(out)


# ---------------- 子命令 ----------------
def cmd_new(args):
    doc = Document()
    if args.title:
        doc.add_heading(args.title, 0)
    doc.save(args.file)
    print(json.dumps({"file": args.file, "title": args.title}, ensure_ascii=False))


def cmd_info(args):
    doc = Document(args.file)
    cp = doc.core_properties
    print(json.dumps({
        "file": args.file,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "title": cp.title,
        "author": cp.author,
    }, ensure_ascii=False, indent=2))


def cmd_read(args):
    doc = Document(args.file)
    blocks = parse_doc(doc)
    if args.max:
        blocks = blocks[:args.max]
    if args.md:
        print(_blocks_to_md(blocks))
    else:
        print(json.dumps({"file": args.file, "blocks": blocks}, ensure_ascii=False, indent=2))


def cmd_search(args):
    doc = Document(args.file)
    needle = args.value.lower()
    hits = []
    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            if needle in item.text.lower():
                hits.append({"type": "paragraph", "text": item.text})
        else:
            for ri, row in enumerate(item.rows, 1):
                for ci, cell in enumerate(row.cells):
                    if needle in cell.text.lower():
                        hits.append({"type": "table", "cell": f"r{ri}c{ci}", "text": cell.text})
    print(json.dumps({"count": len(hits), "hits": hits}, ensure_ascii=False, indent=2))


def _read_structure(args):
    if args.data:
        data = json.loads(args.data) if isinstance(args.data, str) else args.data
    else:
        data = json.load(sys.stdin)
    if not isinstance(data, list):
        raise SystemExit("文档结构必须是 JSON 列表（元素为 heading/paragraph/table/page_break 对象）")
    return data


def cmd_write(args):
    data = _read_structure(args)
    doc = Document(args.file) if os.path.exists(args.file) else Document()
    for item in data:
        t = item.get("type", "paragraph")
        if t == "heading":
            doc.add_heading(item.get("text", ""), level=int(item.get("level", 1)))
        elif t == "table":
            rows = item.get("rows", [])
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    val = row[ci] if ci < len(row) else ""
                    table.cell(ri, ci).text = "" if val is None else str(val)
        elif t == "page_break":
            doc.add_page_break()
        else:  # paragraph
            doc.add_paragraph(item.get("text", ""))
    doc.save(args.file)
    print(json.dumps({"file": args.file, "blocks": len(data)}, ensure_ascii=False))


def main():
    p = argparse.ArgumentParser(prog="word.py", description="Word(.docx) 技能统一入口 CLI")
    sub = p.add_subparsers(dest="cmd", required=True)

    sp = sub.add_parser("new", help="新建文档")
    sp.add_argument("file")
    sp.add_argument("--title")
    sp.set_defaults(func=cmd_new)

    sp = sub.add_parser("info", help="文档元信息")
    sp.add_argument("file")
    sp.set_defaults(func=cmd_info)

    sp = sub.add_parser("read", help="读取段落与表格")
    sp.add_argument("file")
    sp.add_argument("--md", action="store_true", help="输出 Markdown")
    sp.add_argument("--max", type=int, help="最多读取块数")
    sp.set_defaults(func=cmd_read)

    sp = sub.add_parser("search", help="搜索内容")
    sp.add_argument("file")
    sp.add_argument("value")
    sp.set_defaults(func=cmd_search)

    sp = sub.add_parser("write", help="按 JSON 结构写入(新建或追加)")
    sp.add_argument("file")
    sp.add_argument("--data", help="JSON 结构字符串；缺省从 stdin 读取")
    sp.set_defaults(func=cmd_write)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
